"""UI API contracts against explicitly synthetic in-memory SQLite only.

Import isolation prevents app.config/database (and therefore .env or a real
connection) from loading. TestClient is deliberately used without lifespan.
Run: .venv/bin/python -m unittest discover -s tests -p test_ui_api.py -v
"""
from __future__ import annotations

import importlib.util
import json
import re
import sqlite3
import sys
import types
import unittest
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[1]
GROUP = "00000000-0000-4000-8000-000000000001"
OTHER_GROUP = "00000000-0000-4000-8000-000000000002"


def load_isolated_api():
    config = types.ModuleType("app.config")
    config.__dict__.update(HOST="127.0.0.1", PORT=0)
    database = types.ModuleType("app.database")

    def forbidden(*args, **kwargs):
        raise AssertionError("Real database/startup access is forbidden in UI tests")

    database.__dict__.update(get_db=forbidden, init_database=forbidden, ingest_alert_with_dedup=forbidden)
    spec = importlib.util.spec_from_file_location("arbiter_synthetic_ui_api", ROOT / "app/main.py")
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    with patch.dict(sys.modules, {"app.config": config, "app.database": database}):
        spec.loader.exec_module(module)
    return module


class SyntheticDB:
    """Small PostgreSQL-to-SQLite adapter, never a network/database connection.

    SQL predicates, joins, sorting and pagination execute in SQLite rather than
    being replaced with pre-canned responses. Only dialect syntax is translated.
    """

    def __init__(self):
        self.connection = sqlite3.connect(":memory:", check_same_thread=False)
        self.connection.execute("ATTACH DATABASE ':memory:' AS wazuh")
        self.connection.execute("ATTACH DATABASE ':memory:' AS public")
        self.connection.executescript("""
        CREATE TABLE wazuh.alerts (
            id INTEGER PRIMARY KEY, alert_id TEXT, timestamp TEXT, tenant TEXT,
            agent_name TEXT, agent_ip TEXT, rule_id TEXT, rule_level INTEGER,
            description TEXT, rule_groups TEXT, decoder TEXT, location TEXT,
            source_ip TEXT, mitre_ids TEXT, cve TEXT, data TEXT, raw TEXT,
            hit_count INTEGER, first_seen TEXT, last_seen TEXT, wazuh_ids TEXT,
            duplicate_instances TEXT, parsed_details TEXT, is_reviewed INTEGER,
            reviewer_notes TEXT, reviewed_at TEXT, log_signature TEXT
        );
        CREATE TABLE public.ip_intel_cache (
            ip TEXT PRIMARY KEY, abuse_score INTEGER, vt_malicious INTEGER,
            country_code TEXT, isp TEXT, attack_types TEXT
        );
        CREATE TABLE wazuh.analysis_groups (
            id TEXT PRIMARY KEY, title TEXT, category TEXT, analysis TEXT,
            conclusion TEXT, recommendation TEXT, status TEXT, severity TEXT,
            source TEXT, created_at TEXT, updated_at TEXT
        );
        CREATE TABLE wazuh.alert_group_members (
            alert_pk INTEGER, group_id TEXT, UNIQUE(alert_pk, group_id)
        );
        """)
        self.calls = []
        self.add_alert(1, "Alpha", description="Synthetic curl bot", hit_count=3)
        self.add_alert(2, "Beta", description="Synthetic curl bot", hit_count=3)
        self.add_alert(3, "Alpha", description="Synthetic routine", is_reviewed=True)
        self.connection.execute(
            "INSERT INTO wazuh.analysis_groups VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (GROUP, "Synthetic incident", "Test", "Synthetic analysis", "Synthetic conclusion",
             "Synthetic recommendation", "open", "high", "human", "2026-01-01", "2026-01-01"),
        )
        self.connection.executemany("INSERT INTO wazuh.alert_group_members VALUES (?, ?)", [(1, GROUP), (2, GROUP)])
        self.connection.row_factory = self.row

    @staticmethod
    def row(cursor, values):
        record = dict(zip([item[0] for item in cursor.description], values))
        for key, value in record.items():
            if isinstance(value, str) and value.startswith(("[", "{")):
                try:
                    record[key] = json.loads(value)
                except ValueError:
                    pass
        return record

    def add_alert(self, pk, tenant, **changes):
        values = dict(
            id=pk, alert_id=f"synthetic-{pk}", timestamp="2026-01-01T00:00:00Z", tenant=tenant,
            agent_name="synthetic-host", agent_ip="192.0.2.10", rule_id="synthetic-rule", rule_level=7,
            description="Synthetic event", rule_groups="[]", decoder="synthetic", location="synthetic",
            source_ip="192.0.2.1", mitre_ids="[]", cve="", data='{"synthetic":true}',
            raw='{"full_log":"Synthetic curl bot 403 evidence","previous_output":"Synthetic prior"}',
            hit_count=1, first_seen="2026-01-01", last_seen="2026-01-01", wazuh_ids="[]",
            duplicate_instances='[{"synthetic":true}]', parsed_details='{"status":"403","log_type":"http"}',
            is_reviewed=False, reviewer_notes="", log_signature="synthetic-signature",
        )
        values.update(changes)
        columns = ", ".join(values)
        self.connection.execute(
            f"INSERT INTO wazuh.alerts ({columns}) VALUES ({', '.join('?' for _ in values)})", list(values.values())
        )

    def execute(self, sql, params=()):
        self.calls.append((sql, list(params)))
        translated = re.sub(r"\bILIKE\b", "LIKE", sql)
        translated = re.sub(r"::(?:text|uuid|jsonb)", "", translated)
        translated = translated.replace("json_agg(", "json_group_array(").replace("json_build_object(", "json_object(")
        translated = translated.replace("array_agg(", "json_group_array(").replace("'{}'", "'[]'")
        # Expand ANY(list) using the original placeholder positions.
        bind = []
        index = 0

        def placeholder(match):
            nonlocal index
            value = params[index]
            index += 1
            if match.group(0).startswith("ANY"):
                bind.extend(value)
                return "(" + ",".join("?" for _ in value) + ")"
            bind.append(value)
            return "?"

        translated = re.sub(r"ANY\(%s\)|%s", placeholder, translated).replace("= (", "IN (")
        import uuid
        bind = [str(value) if isinstance(value, uuid.UUID) else value for value in bind]
        return self.connection.execute(translated, bind)

    @contextmanager
    def session(self):
        yield self


class UIAPIContractTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.api = load_isolated_api()

    def setUp(self):
        self.db = SyntheticDB()
        self.db_patch = patch.object(self.api, "get_db", self.db.session)
        self.db_patch.start()
        self.client = TestClient(self.api.app)

    def tearDown(self):
        self.client.close()
        self.db_patch.stop()
        self.db.connection.close()

    def test_summary_omits_evidence_payload_without_changing_total(self):
        response = self.client.get("/api/alerts", params={"tenant": "Alpha", "summary": "true", "limit": 1})
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["total"], 2)
        self.assertEqual(len(body["items"]), 1)
        for field in ("raw", "data", "duplicate_instances", "full_log", "previous_output"):
            self.assertNotIn(field, body["items"][0])
        self.assertEqual(body["items"][0]["tenant"], "Alpha")

    def test_group_evidence_retains_tenant_scope(self):
        response = self.client.get("/api/alerts", params={"tenant": "Alpha", "group_id": GROUP})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual([item["id"] for item in response.json()["items"]], [1])
        self.assertEqual(response.json()["total"], 1)

    def test_alert_detail_is_scoped_and_contains_evidence(self):
        response = self.client.get("/api/alerts/1", params={"tenant": "Alpha"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["raw"]["full_log"], "Synthetic curl bot 403 evidence")
        for pk in (2, 999):
            response = self.client.get(f"/api/alerts/{pk}", params={"tenant": "Alpha"})
            self.assertEqual(response.status_code, 404)

    def test_group_metadata_is_scoped_in_list_and_detail(self):
        response = self.client.get("/api/groups", params={"tenant": "Alpha"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(len(response.json()), 1)
        group = response.json()[0]
        self.assertEqual(group["alert_count"], 1)
        self.assertEqual(group["alert_pks"], [1])
        self.assertEqual(group["tenants"], ["Alpha"])
        detail = self.client.get(f"/api/groups/{GROUP}", params={"tenant": "Alpha"})
        self.assertEqual(detail.status_code, 200, detail.text)
        self.assertEqual(detail.json(), group)
        self.assertEqual(detail.json()["analysis"], "Synthetic analysis")
        self.assertEqual(self.client.get(f"/api/groups/{GROUP}", params={"tenant": "Gamma"}).status_code, 404)
        self.assertEqual(self.client.get(f"/api/groups/{OTHER_GROUP}").status_code, 404)
        self.assertEqual(self.client.get("/api/groups", params={"tenant": "Gamma"}).json(), [])

    def test_suggestions_use_the_active_triage_scope(self):
        self.db.add_alert(4, "Alpha", description="Synthetic curl bot", hit_count=4)
        self.db.connection.execute("INSERT INTO wazuh.alert_group_members VALUES (?, ?)", (4, GROUP))
        response = self.client.get("/api/suggestions", params={
            "tenant": "Alpha", "q": "curl", "review_status": "unreviewed",
            "filter_preset": "repeating", "status": "grouped", "group_id": GROUP,
        })
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["mode"], "local_rules")
        self.assertIsNone(body["warning"])
        self.assertTrue(body["suggestions"])
        ids = {pk for suggestion in body["suggestions"] for pk in suggestion["alert_ids"]}
        self.assertEqual(ids, {"1", "4"})

    def test_preview_matches_all_triage_filters_without_losing_group_tenant(self):
        params = {"tenant": "Alpha", "q": "curl", "review_status": "unreviewed",
                  "filter_preset": "repeating", "status": "grouped", "group_id": GROUP}
        response = self.client.get("/api/export/preview", params=params)
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["count"], 1)
        self.assertEqual(body["meta"]["title"], "Synthetic incident")
        self.assertEqual(body["scope"]["tenant"], "Alpha")
        self.assertFalse(body["truncated"])
        self.assertIn("Synthetic curl bot", body["html"])
        self.assertNotIn("Synthetic routine", body["html"])
        self.assertIn("<strong>Total Alert:</strong> 1 grup unik", body["html"])
        params["q"] = "no-synthetic-match"
        self.assertEqual(self.client.get("/api/export/preview", params=params).json()["count"], 0)

    def test_every_export_uses_the_preview_scope_and_count(self):
        import csv
        import io
        from docx import Document
        from openpyxl import load_workbook

        params = {"tenant": "Alpha", "q": "bot", "review_status": "unreviewed",
                  "filter_preset": "repeating", "status": "grouped", "group_id": GROUP}
        preview = self.client.get("/api/export/preview", params=params).json()
        routes = ["/api/export/report-html", "/api/export/docx", "/api/export/xlsx",
                  "/api/export/tsv", "/api/export/markdown", "/api/export.csv"]
        for route in routes:
            with self.subTest(route=route):
                response = self.client.get(route, params=params)
                self.assertEqual(response.status_code, 200, response.text[:80])
                self.assertEqual(response.headers.get("X-Arbiter-Count"), str(preview["count"]))
                if route.endswith("report-html"):
                    self.assertEqual(response.text, preview["html"])
                elif route.endswith("docx"):
                    document = Document(io.BytesIO(response.content))
                    self.assertIn(preview["meta"]["title"], " ".join(p.text for p in document.paragraphs))
                    self.assertEqual(len(document.tables[0].rows), 2)
                elif route.endswith("xlsx"):
                    sheet = load_workbook(io.BytesIO(response.content)).active
                    assert sheet is not None
                    self.assertEqual(sheet.max_row, 2)
                elif route.endswith("tsv"):
                    self.assertEqual(len(list(csv.reader(io.StringIO(response.text), delimiter="\t"))), 2)
                elif route.endswith(".csv"):
                    rows = list(csv.DictReader(io.StringIO(response.text)))
                    self.assertEqual([row["alert_id"] for row in rows], ["synthetic-1"])
                else:
                    self.assertIn(preview["meta"]["title"], response.text)

    def test_export_bound_is_explicit_and_never_silently_truncates(self):
        routes = ["/api/export/preview", "/api/export/report-html", "/api/export/docx",
                  "/api/export/xlsx", "/api/export/tsv", "/api/export/markdown", "/api/export.csv"]
        with patch.object(self.api, "MAX_EXPORT_ALERTS", 2, create=True):
            for route in routes:
                with self.subTest(route=route):
                    response = self.client.get(route)
                    self.assertEqual(response.status_code, 413, route)
                    self.assertIn("2", response.json()["detail"])
            response = self.client.get("/api/export/preview", params={"tenant": "Alpha"})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["count"], 2)

    def test_selected_export_overrides_filters_but_never_tenant(self):
        response = self.client.get("/api/export/preview", params={
            "tenant": "Alpha", "alert_ids": "1,3", "q": "no-match",
            "review_status": "unreviewed", "filter_preset": "repeating",
            "status": "ungrouped", "group_id": OTHER_GROUP,
        })
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["count"], 2)
        self.assertIsNone(body["scope"]["group_id"])
        self.assertEqual(body["scope"]["q"], "")
        self.assertEqual(body["scope"]["tenant"], "Alpha")
        self.assertEqual(body["meta"]["title"], "Investigasi Kasus (2 Alert Terpilih)")
        response = self.client.get("/api/export.csv", params={"tenant": "Alpha", "alert_ids": "1,3"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("synthetic-1", response.text)
        self.assertIn("synthetic-3", response.text)
        self.assertNotIn("synthetic-2", response.text)

    def test_invalid_explicit_export_ids_never_broaden_or_partially_match(self):
        for ids in ("", "garbage", "1,garbage", "-1", "0", "1,", "1.0", "1,,3"):
            with self.subTest(ids=ids):
                response = self.client.get("/api/export/preview", params={"alert_ids": ids})
                self.assertEqual(response.status_code, 422, response.text[:80])
        response = self.client.get("/api/export/preview", params={"alert_ids": "1,1,3", "tenant": "Alpha"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["count"], 2)
        self.assertEqual(response.json()["scope"]["alert_ids"], "1,3")

    def test_membership_add_rejects_the_entire_cross_tenant_selection(self):
        response = self.client.post(f"/api/groups/{GROUP}/members?tenant=Alpha", json={"alert_pks": [3, 2]})
        self.assertEqual(response.status_code, 400, response.text)
        members = self.db.connection.execute(
            "SELECT alert_pk FROM wazuh.alert_group_members WHERE group_id = ? ORDER BY alert_pk", (GROUP,)
        ).fetchall()
        self.assertEqual([row["alert_pk"] for row in members], [1, 2])
        response = self.client.post(f"/api/groups/{GROUP}/members?tenant=Alpha", json={"alert_pks": [3]})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.client.get(f"/api/groups/{GROUP}?tenant=Alpha").json()["alert_count"], 2)

    def test_group_create_validates_tenant_before_any_insert(self):
        response = self.client.post("/api/groups?tenant=Alpha", json={"title": "Synthetic new", "alert_pks": [1, 2]})
        self.assertEqual(response.status_code, 400, response.text)
        self.assertEqual(self.db.connection.execute("SELECT count(*) AS n FROM wazuh.analysis_groups").fetchone()["n"], 1)
        self.assertFalse(any("INSERT" in sql for sql, _ in self.db.calls))
        response = self.client.post("/api/groups?tenant=Alpha", json={"title": "Synthetic new", "alert_pks": [1, 3]})
        self.assertEqual(response.status_code, 200, response.text)
        detail = self.client.get(f"/api/groups/{response.json()['id']}?tenant=Alpha")
        self.assertEqual(detail.status_code, 200, detail.text)
        self.assertEqual(detail.json()["alert_count"], 2)

    def test_tied_sort_pagination_has_a_stable_id_tiebreaker(self):
        self.db.connection.execute("UPDATE wazuh.alerts SET hit_count = 1, rule_level = 7")
        for sort_by in ("hits", "severity", "timestamp"):
            with self.subTest(sort_by=sort_by):
                ids = []
                for offset in range(3):
                    response = self.client.get("/api/alerts", params={"sort_by": sort_by, "limit": 1, "offset": offset})
                    self.assertEqual(response.status_code, 200, response.text)
                    ids.extend(item["id"] for item in response.json()["items"])
                self.assertEqual(ids, [3, 2, 1])

    def test_static_mount_serves_only_its_directory_without_database(self):
        import tempfile

        mount = next((route for route in self.api.app.routes if route.path == "/static"), None)
        self.assertIsNotNone(mount, "The /static mount is required for native ES modules and CSS")
        assert mount is not None
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "synthetic.js").write_text("export const synthetic = true;", encoding="utf-8")
            with patch.object(mount.app, "directory", temporary), patch.object(mount.app, "all_directories", [temporary]):
                response = self.client.get("/static/synthetic.js")
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.text, "export const synthetic = true;")
                self.assertEqual(self.client.get("/static/missing.js").status_code, 404)
                self.assertEqual(self.client.get("/static/%2e%2e/main.py").status_code, 404)
        self.assertEqual(self.db.calls, [])

    def test_membership_remove_refuses_out_of_scope_evidence(self):
        response = self.client.delete(f"/api/groups/{GROUP}/members/2?tenant=Alpha")
        self.assertEqual(response.status_code, 400, response.text)
        self.assertEqual(self.client.get(f"/api/groups/{GROUP}?tenant=Beta").json()["alert_count"], 1)
        response = self.client.delete(f"/api/groups/{GROUP}/members/1?tenant=Alpha")
        self.assertEqual(response.status_code, 409, response.text)
        self.assertEqual(self.client.get(f"/api/groups/{GROUP}?tenant=Alpha").status_code, 200)

    def test_malformed_group_ids_are_rejected_before_database_execution(self):
        for route, params in (
            ("/api/alerts", {"group_id": "not-a-uuid"}),
            ("/api/export/preview", {"group_id": "not-a-uuid"}),
            ("/api/suggestions", {"group_id": "not-a-uuid"}),
            ("/api/groups/not-a-uuid", {}),
        ):
            with self.subTest(route=route):
                response = self.client.get(route, params=params)
                self.assertEqual(response.status_code, 422, response.text[:80])
        self.assertEqual(self.db.calls, [])

    def test_pagination_rejects_negative_offset_before_querying(self):
        response = self.client.get("/api/alerts", params={"offset": -1})
        self.assertEqual(response.status_code, 422, response.text[:80])
        self.assertEqual(self.db.calls, [])

    def test_database_status_endpoint_returns_profiles_and_telemetry(self):
        response = self.client.get("/api/database/status")
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()
        self.assertEqual(data["status"], "ok")
        self.assertIn("active_profile", data)
        self.assertIn("profiles", data)
        self.assertIn("stats", data)
        self.assertGreaterEqual(data["stats"]["alerts"], 2)

    def test_alerts_support_time_range_filters(self):
        response = self.client.get("/api/alerts", params={"tenant": "Alpha", "time_from": "2026-01-01T00:00:00Z"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(len(response.json()["items"]), 2)
        response_future = self.client.get("/api/alerts", params={"tenant": "Alpha", "time_from": "2099-01-01T00:00:00Z"})
        self.assertEqual(response_future.status_code, 200)
        self.assertEqual(len(response_future.json()["items"]), 0)

    def test_intel_enrich_endpoint_private_ip(self):
        response = self.client.post("/api/intel/enrich", json={"ips": ["10.0.0.1", "192.168.1.100"]})
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()
        self.assertEqual(data["status"], "ok")
        self.assertEqual(data["count"], 2)
        for item in data["enriched"]:
            self.assertEqual(item["status"], "private")
            self.assertEqual(item["abuse_score"], 0)
            self.assertEqual(item["country_code"], "LAN")

    def test_intel_enrich_endpoint_rejects_empty_or_invalid(self):
        response = self.client.post("/api/intel/enrich", json={"ips": ["invalid-ip", "not_an_ip"]})
        self.assertEqual(response.status_code, 400)


if __name__ == "__main__":
    unittest.main()
