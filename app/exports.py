"""Deliverables and report export engines for Arbiter.

Supports:
1. GDocs Table Report (Semantic HTML designed for direct Ctrl+V clipboard paste into Google Docs)
2. Word Document (.docx landscape with formal SOC headers and tables)
3. Excel Workbook (.xlsx styled sheets with threat intelligence metrics)
4. Tab-Separated Values (TSV clipboard text)
5. Markdown Executive Briefing
"""

from __future__ import annotations

import csv
import io
from html import escape
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def generate_tsv_table(alerts: list[dict[str, Any]]) -> str:
    """Generates a clean Tab-Separated Values table for instant pasting into Excel or Word."""
    headers = [
        "No",
        "Hit Count",
        "Timestamp",
        "Tenant",
        "Agent / Host",
        "Rule ID",
        "Level",
        "Event Description",
        "Source IP",
        "Abuse Score (%)",
        "VT Detections",
        "Country",
        "Full Log Snippet",
        "Status",
    ]
    rows = [headers]
    for idx, a in enumerate(alerts, 1):
        full_log = a.get("full_log") or (a.get("raw") or {}).get("full_log") or ""
        clean_log = str(full_log).replace("\t", " ").replace("\r\n", " ").replace("\n", " ").strip()
        reviewed = "Reviewed" if a.get("is_reviewed") else "Pending"
        rows.append([
            str(idx),
            str(a.get("hit_count") or 1),
            str(a.get("timestamp") or "-"),
            str(a.get("tenant") or "Default"),
            str(a.get("agent_name") or "-"),
            str(a.get("rule_id") or "-"),
            str(a.get("rule_level") or 0),
            str(a.get("description") or "-").replace("\t", " "),
            str(a.get("source_ip") or "-"),
            str(a.get("abuse_score") if a.get("abuse_score") is not None else "-"),
            str(a.get("vt_malicious") if a.get("vt_malicious") is not None else "-"),
            str(a.get("country_code") or "-"),
            clean_log,
            reviewed,
        ])
    output = io.StringIO()
    writer = csv.writer(output, delimiter="\t", lineterminator="\n")
    writer.writerows(rows)
    return output.getvalue()


def generate_markdown_report(alerts: list[dict[str, Any]], meta: dict[str, Any]) -> str:
    """Generates an executive SOC investigation summary in Markdown format."""
    total_hits = sum(int(a.get("hit_count") or 1) for a in alerts)
    repeated = len([a for a in alerts if int(a.get("hit_count") or 1) > 1])
    highest_lvl = max([int(a.get("rule_level") or 0) for a in alerts], default=0)

    title = meta.get("title") or "SOC Security Telemetry Report"
    category = meta.get("category") or "General Incident"
    severity = meta.get("severity") or ("High" if highest_lvl >= 10 else "Medium")
    status = meta.get("status") or "Investigating"

    lines = [
        f"# {title}",
        f"**Kategori:** {category} | **Severity:** {severity.upper()} | **Status:** {status.upper()}",
        f"- **Total Alert Unik:** {len(alerts)}",
        f"- **Total Kemunculan (Hits):** {total_hits} (dengan {repeated} kluster berulang)",
        f"- **Tingkat Severity Maksimum:** Level {highest_lvl}",
        "",
        "## 1. Rekapitulasi Kronologi Log & Threat Intelligence",
    ]

    for idx, a in enumerate(alerts, 1):
        ip = a.get("source_ip") or "-"
        intel_notes = []
        if a.get("abuse_score") is not None:
            intel_notes.append(f"Abuse: {a['abuse_score']}%")
        if a.get("vt_malicious"):
            intel_notes.append(f"VT: {a['vt_malicious']}")
        if a.get("country_code"):
            intel_notes.append(str(a["country_code"]))
        intel_str = f" [{', '.join(intel_notes)}]" if intel_notes else ""

        hits_str = f" (x{a['hit_count']})" if int(a.get("hit_count") or 1) > 1 else ""
        lines.append(
            f"{idx}. `[{a.get('timestamp') or '-'}]` **Level {a.get('rule_level') or 0}** (Rule {a.get('rule_id')}) - "
            f"{a.get('description')}{hits_str} | IP: `{ip}`{intel_str} | Host: `{a.get('agent_name') or '-'}`"
        )

    lines.extend([
        "",
        "## 2. Analisis Teknis Investigasi",
        meta.get("analysis") or "Aktivitas serangan terekam secara kronologis pada tabel di atas. Diperlukan mitigasi berkelanjutan.",
        "",
        "## 3. Kesimpulan Investigasi (Verdict)",
        meta.get("conclusion") or "True Positive — Terkonfirmasi adanya anomali atau upaya akses yang terdeteksi Wazuh SIEM.",
        "",
        "## 4. Rekomendasi Tindakan (Containment & Mitigation)",
        meta.get("recommendation") or "1. Blokir alamat IP penyerang pada edge firewall / perimeter.\n2. Verifikasi status keamanan endpoint terkait.\n3. Lakukan audit pada rule deteksi.",
    ])
    return "\n".join(lines)


def generate_gdocs_html_report(alerts: list[dict[str, Any]], meta: dict[str, Any]) -> str:
    """Generates styled semantic HTML with clean table borders and colors for Google Docs pasting."""
    title = meta.get("title") or "Laporan Hasil Investigasi Insiden Keamanan (SOC Incident Report)"
    category = meta.get("category") or "Security Incident"
    severity = meta.get("severity") or "Medium"
    status = meta.get("status") or "Open"
    analysis = meta.get("analysis") or ""
    conclusion = meta.get("conclusion") or ""
    recommendation = meta.get("recommendation") or ""

    lvl_class = "badge-high" if severity.lower() in ("high", "critical") else ("badge-med" if severity.lower() == "medium" else "badge-low")

    unique_ips = list(dict.fromkeys(a["source_ip"] for a in alerts if a.get("source_ip")))
    unique_agents = list(dict.fromkeys(a["agent_name"] for a in alerts if a.get("agent_name")))
    first_time = alerts[0].get("timestamp") if alerts else "-"
    last_time = alerts[-1].get("timestamp") if alerts else "-"
    total_hits = sum(int(a.get("hit_count") or 1) for a in alerts)

    html_parts = [
        """<!DOCTYPE html><html><head><meta charset="utf-8">
        <style>
        body { font-family: Arial, sans-serif; font-size: 11pt; color: #111; line-height: 1.45; }
        h1 { color: #1a365d; border-bottom: 2px solid #2b6cb0; padding-bottom: 6px; font-size: 16pt; margin-top: 10px; }
        h2 { color: #2c5282; font-size: 13pt; margin-top: 20px; border-bottom: 1px solid #cbd5e0; padding-bottom: 4px; }
        table { width: 100%; border-collapse: collapse; margin: 12px 0 20px; font-size: 9.5pt; }
        th, td { border: 1px solid #cbd5e0; padding: 7px 10px; text-align: left; vertical-align: top; }
        th { background-color: #edf2f7; color: #2d3748; font-weight: bold; }
        tr:nth-child(even) { background-color: #f7fafc; }
        .badge-high { background-color: #fed7d7; color: #9b2c2c; padding: 2px 6px; border-radius: 4px; font-weight: bold; }
        .badge-med { background-color: #feebc8; color: #7b341e; padding: 2px 6px; border-radius: 4px; font-weight: bold; }
        .badge-low { background-color: #e2e8f0; color: #4a5568; padding: 2px 6px; border-radius: 4px; }
        .meta-box { background: #f8fafc; border-left: 4px solid #3182ce; padding: 12px; margin: 12px 0; font-size: 10pt; }
        .content-box { background: #ffffff; border: 1px solid #cbd5e0; padding: 12px; border-radius: 4px; margin-bottom: 16px; font-size: 10.5pt; }
        </style></head><body>
        """,
        f"<h1>{escape(title)}</h1>",
        f"""<div class="meta-box">
            <p style="margin: 2px 0;"><strong>Kategori:</strong> {escape(category)} | <strong>Severity:</strong> <span class="{lvl_class}">{escape(severity.upper())}</span> | <strong>Status:</strong> {escape(status.upper())}</p>
            <p style="margin: 2px 0;"><strong>Rentang Waktu:</strong> {escape(str(first_time))} s/d {escape(str(last_time))}</p>
            <p style="margin: 2px 0;"><strong>Total Alert:</strong> {len(alerts)} grup unik ({total_hits} total hits) | <strong>Host Terdampak:</strong> {escape(', '.join(unique_agents) or '-')} | <strong>IP Penyerang:</strong> {escape(', '.join(unique_ips) or '-')}</p>
        </div>
        <h2>1. Tabel Kronologi Investigasi (Attack Chronology)</h2>
        """,
    ]

    if alerts:
        html_parts.append("""
        <table border="1" cellpadding="8" cellspacing="0" style="width: 100%; border-collapse: collapse; border: 1px solid #cbd5e0; margin: 12px 0 20px; font-family: Arial, sans-serif; font-size: 9.5pt;">
            <thead>
                <tr style="background-color: #edf2f7;">
                    <th style="width: 25px; border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top; font-weight: bold; color: #2d3748;">No</th>
                    <th style="width: 45px; border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top; font-weight: bold; color: #2d3748;">Hits</th>
                    <th style="width: 135px; border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top; font-weight: bold; color: #2d3748;">Timestamp</th>
                    <th style="width: 120px; border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top; font-weight: bold; color: #2d3748;">Host / Agent</th>
                    <th style="width: 110px; border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top; font-weight: bold; color: #2d3748;">Rule &amp; Severity</th>
                    <th style="border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top; font-weight: bold; color: #2d3748;">Ringkasan Event / Full Log</th>
                    <th style="width: 170px; border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top; font-weight: bold; color: #2d3748;">Source IP &amp; Intel</th>
                    <th style="width: 120px; border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top; font-weight: bold; color: #2d3748;">CVE / MITRE</th>
                </tr>
            </thead>
            <tbody>
        """)
        for idx, a in enumerate(alerts, 1):
            lvl = int(a.get("rule_level") or 0)
            row_lvl = "badge-high" if lvl >= 10 else ("badge-med" if lvl >= 7 else "badge-low")
            hits = int(a.get("hit_count") or 1)

            intel_txt = ""
            if a.get("abuse_score") is not None:
                intel_txt = f"<br><small><strong>Abuse:</strong> {a['abuse_score']}%"
                if a.get("vt_malicious"):
                    intel_txt += f" | <strong>VT:</strong> {a['vt_malicious']}"
                if a.get("country_code"):
                    intel_txt += f" ({escape(str(a['country_code']))})"
                intel_txt += "</small>"

            cve_mitre = []
            if a.get("cve"):
                cve_mitre.append(f"<strong>CVE:</strong> {escape(str(a['cve']))}")
            if a.get("mitre_ids"):
                cve_mitre.append(f"<strong>MITRE:</strong> {escape(', '.join(a['mitre_ids']))}")

            full_log = a.get("full_log") or (a.get("raw") or {}).get("full_log") or ""
            log_block = (
                f"<div style='margin-top: 6px; padding: 6px 8px; background-color: #f8fafc; "
                f"border: 1px solid #e2e8f0; border-radius: 4px; font-family: Consolas, \"Courier New\", monospace; "
                f"font-size: 8pt; color: #2d3748; white-space: pre-wrap; word-break: break-all;'>"
                f"<strong>Full Log:</strong> {escape(str(full_log))}</div>"
            ) if full_log else ""

            row_bg = "background-color: #f7fafc;" if idx % 2 == 0 else "background-color: #ffffff;"
            html_parts.append(f"""
                <tr style="{row_bg}">
                    <td style="border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top;">{idx}</td>
                    <td style="border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top;"><strong>{hits}x</strong></td>
                    <td style="border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top;">{escape(str(a.get('timestamp') or '-'))}</td>
                    <td style="border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top;"><strong>{escape(str(a.get('agent_name') or '-'))}</strong><br><small>{escape(str(a.get('agent_ip') or '-'))}</small></td>
                    <td style="border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top;"><span class="{row_lvl}">Level {lvl}</span><br><small>Rule {escape(str(a.get('rule_id') or '-'))}</small></td>
                    <td style="border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top;"><strong>{escape(str(a.get('description') or '-'))}</strong>{log_block}</td>
                    <td style="border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top;">{escape(str(a.get('source_ip') or '-'))}{intel_txt}</td>
                    <td style="border: 1px solid #cbd5e0; padding: 8px 10px; text-align: left; vertical-align: top;">{'<br>'.join(cve_mitre) or '-'}</td>
                </tr>
            """)
        html_parts.append("</tbody></table>")
    else:
        html_parts.append("<p><em>Tidak ada alert terkait dalam laporan ini.</em></p>")

    html_parts.append(f"""
    <h2>2. Kronologi & Analisis Investigasi</h2>
    <div class="content-box">
        {escape(analysis).replace(chr(10), '<br>') or '<em>Aktivitas serangan terekam secara kronologis pada tabel di atas. Diperlukan penelusuran lebih lanjut terkait payload yang dieksekusi.</em>'}
    </div>

    <h2>3. Kesimpulan Investigasi (Verdict)</h2>
    <div class="content-box">
        {escape(conclusion).replace(chr(10), '<br>') or '<em>True Positive — Terkonfirmasi adanya aktivitas anomali / serangan yang terdeteksi oleh Wazuh SIEM.</em>'}
    </div>

    <h2>4. Rekomendasi Tindakan (Containment & Remediation)</h2>
    <div class="content-box">
        {escape(recommendation).replace(chr(10), '<br>') or '<em>1. Terapkan pemblokiran IP penyerang pada Edge Firewall / WAF.<br>2. Verifikasi integritas akun dan file pada host terdampak.<br>3. Pastikan security patches dan konfigurasi hardening telah diperbarui.</em>'}
    </div>
    </body></html>
    """)
    return "".join(html_parts)


def generate_docx_landscape_report(alerts: list[dict[str, Any]], meta: dict[str, Any]) -> io.BytesIO:
    """Generates an official Word document (.docx) in landscape orientation."""
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    # Swap page dimensions for landscape
    new_width, new_height = section.page_height, section.page_width
    section.page_width = new_width
    section.page_height = new_height
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    title = meta.get("title") or "SOC Security Telemetry & Incident Report"
    category = meta.get("category") or "Security Incident"
    severity = meta.get("severity") or "Medium"
    status = meta.get("status") or "Open"

    title_p = doc.add_paragraph()
    title_run = title_p.add_run(title)
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(26, 54, 93)

    meta_p = doc.add_paragraph()
    meta_p.add_run(f"Kategori: {category}  |  Severity: {severity.upper()}  |  Status: {status.upper()}\n").bold = True
    first_time = alerts[0].get("timestamp") if alerts else "-"
    last_time = alerts[-1].get("timestamp") if alerts else "-"
    meta_p.add_run(f"Rentang Waktu: {first_time} s/d {last_time}  |  Total Alert: {len(alerts)}")

    doc.add_heading("1. Tabel Kronologi Investigasi", level=2)

    headers = ["No", "Hits", "Waktu", "Host", "Rule ID & Lvl", "Event Description", "Source IP & Intel", "CVE / MITRE"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(9.5)

    # Data Rows
    for idx, a in enumerate(alerts, 1):
        row_cells = table.add_row().cells
        hits = str(a.get("hit_count") or 1)
        intel = []
        if a.get("abuse_score") is not None:
            intel.append(f"Abuse: {a['abuse_score']}%")
        if a.get("country_code"):
            intel.append(str(a["country_code"]))
        intel_str = f" ({', '.join(intel)})" if intel else ""

        row_cells[0].text = str(idx)
        row_cells[1].text = f"{hits}x"
        row_cells[2].text = str(a.get("timestamp") or "-")
        row_cells[3].text = str(a.get("agent_name") or "-")
        row_cells[4].text = f"Rule {a.get('rule_id')} (Lv {a.get('rule_level')})"
        row_cells[5].text = str(a.get("description") or "-")
        row_cells[6].text = f"{a.get('source_ip') or '-'}{intel_str}"
        row_cells[7].text = str(a.get("cve") or ", ".join(a.get("mitre_ids") or []) or "-")

        for c in row_cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)

    doc.add_heading("2. Analisis Investigasi", level=2)
    doc.add_paragraph(meta.get("analysis") or "Aktivitas serangan terekam secara kronologis pada tabel di atas.")

    doc.add_heading("3. Kesimpulan Investigasi (Verdict)", level=2)
    doc.add_paragraph(meta.get("conclusion") or "True Positive — Terkonfirmasi adanya aktivitas serangan.")

    doc.add_heading("4. Rekomendasi Tindakan", level=2)
    doc.add_paragraph(meta.get("recommendation") or "1. Blokir IP penyerang pada firewall.\n2. Verifikasi host target.")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def generate_xlsx_workbook(alerts: list[dict[str, Any]]) -> io.BytesIO:
    """Generates a styled Excel workbook (.xlsx) with threat intelligence metrics."""
    wb = Workbook()
    ws = wb.worksheets[0]
    ws.title = "Wazuh Alerts Triage"

    headers = [
        "No",
        "Hit Count",
        "Timestamp",
        "Tenant",
        "Agent Name",
        "Agent IP",
        "Rule ID",
        "Level",
        "Event Description",
        "Source IP",
        "Abuse Score (%)",
        "VT Detections",
        "Country",
        "Location",
        "Decoder",
        "CVE",
        "MITRE ATT&CK",
        "Review Status",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E0"),
        right=Side(style="thin", color="CBD5E0"),
        top=Side(style="thin", color="CBD5E0"),
        bottom=Side(style="thin", color="CBD5E0"),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for idx, a in enumerate(alerts, 1):
        row_vals = [
            idx,
            int(a.get("hit_count") or 1),
            a.get("timestamp") or "",
            a.get("tenant") or "Default",
            a.get("agent_name") or "",
            a.get("agent_ip") or "",
            a.get("rule_id") or "",
            int(a.get("rule_level") or 0),
            a.get("description") or "",
            a.get("source_ip") or "",
            a.get("abuse_score") if a.get("abuse_score") is not None else "",
            a.get("vt_malicious") if a.get("vt_malicious") is not None else "",
            a.get("country_code") or "",
            a.get("location") or "",
            a.get("decoder") or "",
            a.get("cve") or "",
            ", ".join(a.get("mitre_ids") or []),
            "Reviewed" if a.get("is_reviewed") else "Pending",
        ]
        ws.append(row_vals)
        for col_idx in range(1, len(row_vals) + 1):
            cell = ws.cell(row=idx + 1, column=col_idx)
            cell.font = Font(name="Arial", size=9)
            cell.border = thin_border

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
